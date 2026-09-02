"""
Exporta o estado atual do banco (app regras) para um arquivo referencias.xlsx
no layout exato que config.py/o motor de cálculo esperam. Usado:
  1. Por apps/calculo/services.py — materializa as regras atuais antes de
     chamar main.py via subprocess (fluxo normal de cálculo).
  2. Para backup/auditoria point-in-time.
  3. Para dar à equipe de RH uma visão "tudo numa planilha só" quando precisar.

Uso:
  python manage.py exportar_referencias --saida referencias_export.xlsx
"""

from pathlib import Path

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.regras.models import (
    Cargo, Grupo, IndicadorRegra, Loja, MetaFaturamentoRegional,
    MetaMensal, ParametroTOTVS,
)


def _estilizar_cabecalho(ws, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 20


def _escrever(ws, cabecalho, linhas):
    ws.append(cabecalho)
    _estilizar_cabecalho(ws, len(cabecalho))
    for linha in linhas:
        ws.append(list(linha))


def _num(valor):
    return float(valor) if valor is not None else None


class Command(BaseCommand):
    help = "Exporta o banco de regras para um referencias.xlsx no layout esperado pelo motor de cálculo."

    def add_arguments(self, parser):
        parser.add_argument("--saida", default="referencias_export.xlsx", help="Caminho do arquivo a gerar")

    def handle(self, *args, **options):
        caminho = Path(options["saida"])
        wb = Workbook()
        wb.remove(wb.active)

        self._escrever_grupos(wb)
        self._escrever_cargos(wb)
        self._escrever_regras_calculo(wb)
        self._escrever_tetos(wb)
        self._escrever_metas_mensais(wb)
        self._escrever_metas_faturamento_regional(wb)
        self._escrever_lojas(wb)
        self._escrever_totvs(wb)

        wb.save(caminho)
        self.stdout.write(self.style.SUCCESS(f"Exportado: {caminho.resolve()}"))

    def _escrever_grupos(self, wb):
        ws = wb.create_sheet("Grupos")
        linhas = [(g.nome, g.modelo) for g in Grupo.objects.order_by("nome")]
        _escrever(ws, ["grupo", "modelo"], linhas)

    def _escrever_cargos(self, wb):
        ws = wb.create_sheet("Cargos")
        linhas = [(c.nome, c.grupo.nome) for c in Cargo.objects.select_related("grupo").order_by("nome")]
        _escrever(ws, ["cargo", "grupo_calculo"], linhas)

    def _escrever_regras_calculo(self, wb):
        ws = wb.create_sheet("Regras_Calculo")
        linhas = []
        indicadores = (
            IndicadorRegra.objects
            .select_related("grupo", "loja")
            .prefetch_related("faixas")
            .order_by("grupo__nome", "indicador", "loja__unidade")
        )
        for ind in indicadores:
            for faixa in ind.faixas.order_by("faixa"):
                linhas.append((
                    ind.grupo.nome, ind.indicador, faixa.faixa,
                    _num(faixa.valor_min), _num(faixa.valor_max), _num(faixa.pct),
                    ind.direcao, ind.chave_meta or None,
                    ind.depende_indicador or None, _num(ind.depende_valor_min),
                    ind.loja.unidade if ind.loja else None,
                ))
        _escrever(
            ws,
            ["grupo", "indicador", "faixa", "valor_min", "valor_max", "pct", "direcao", "chave_meta",
             "depende_indicador", "depende_valor_min", "unidade"],
            linhas,
        )

    def _escrever_tetos(self, wb):
        ws = wb.create_sheet("Tetos")
        linhas = [
            (c.nome, _num(c.teto))
            for c in Cargo.objects.filter(teto__isnull=False).order_by("nome")
        ]
        _escrever(ws, ["cargo", "teto"], linhas)

    def _escrever_metas_mensais(self, wb):
        ws = wb.create_sheet("Metas_Mensais")
        linhas = [
            (m.mes, m.chave, m.faixa, _num(m.valor_min), _num(m.valor_max))
            for m in MetaMensal.objects.order_by("mes", "chave", "faixa")
        ]
        _escrever(ws, ["mes", "chave", "faixa", "valor_min", "valor_max"], linhas)

    def _escrever_metas_faturamento_regional(self, wb):
        ws = wb.create_sheet("Metas_Faturamento_Regional")
        linhas = [
            (m.mes, m.regional, _num(m.meta))
            for m in MetaFaturamentoRegional.objects.order_by("mes", "regional")
        ]
        _escrever(ws, ["mes", "regional", "meta"], linhas)

    def _escrever_lojas(self, wb):
        ws = wb.create_sheet("Lojas")
        linhas = [(l.unidade, l.regional) for l in Loja.objects.order_by("regional", "unidade")]
        _escrever(ws, ["unidade", "regional"], linhas)

    def _escrever_totvs(self, wb):
        ws = wb.create_sheet("TOTVS")
        linhas = [(p.parametro, p.valor, p.descricao) for p in ParametroTOTVS.objects.order_by("parametro")]
        _escrever(ws, ["parametro", "valor", "descricao"], linhas)
